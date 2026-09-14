"""app/services/reports/export_service.py — every exported file embeds the
"RELATÓRIO AUXILIAR" disclaimer and never claims to be an official TSE
filing (absolute rule of this project's scope)."""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import openpyxl
from pypdf import PdfReader

from app.models.enums import ExpenseStatus
from app.services.reports.export_service import DISCLAIMER, export_csv, export_pdf, export_xlsx

_COLUMNS = [("date", "Data"), ("supplier_name", "Fornecedor"), ("amount", "Valor (R$)"), ("status", "Status")]

_ROWS = [
    {
        "date": date(2024, 3, 15),
        "supplier_name": "Gráfica São José Ltda",
        "amount": Decimal("1234.56"),
        "status": ExpenseStatus.COMPLETE,
    },
    {"date": None, "supplier_name": "Fornecedor não informado", "amount": Decimal("99.90"), "status": None},
]


def test_csv_contains_disclaimer_and_all_rows():
    content = export_csv(_ROWS, columns=_COLUMNS, title="Despesas").decode("utf-8-sig")
    assert DISCLAIMER in content
    assert "Gráfica São José Ltda" in content
    assert "1234.56" in content
    assert "COMPLETE" in content
    assert "Fornecedor não informado" in content


def test_csv_never_claims_official_tse_submission():
    content = export_csv(_ROWS, columns=_COLUMNS, title="Despesas").decode("utf-8-sig")
    assert "NÃO é uma prestação de contas oficial" in content
    assert "CONTA+JE" in content


def test_xlsx_round_trips_decimal_and_date_and_enum():
    content = export_xlsx(_ROWS, columns=_COLUMNS, title="Despesas")
    workbook = openpyxl.load_workbook(io.BytesIO(content))
    sheet = workbook.active

    all_text = "\n".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)
    assert DISCLAIMER in all_text
    assert "Gráfica São José Ltda" in all_text

    data_rows = list(sheet.iter_rows(values_only=True))
    amount_row = next(r for r in data_rows if r[1] == "Gráfica São José Ltda")
    assert float(amount_row[2]) == 1234.56
    assert amount_row[0].date() == date(2024, 3, 15)


def test_pdf_has_valid_header_and_embeds_disclaimer_text():
    content = export_pdf(_ROWS, columns=_COLUMNS, title="Despesas")
    assert content[:5] == b"%PDF-"

    reader = PdfReader(io.BytesIO(content))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "RELATÓRIO AUXILIAR" in text
    assert "CONTA+JE" in text
    assert "Gráfica São José Ltda" in text


def test_export_with_no_rows_still_produces_a_valid_file_with_headers():
    csv_content = export_csv([], columns=_COLUMNS, title="Despesas").decode("utf-8-sig")
    assert "Data,Fornecedor" in csv_content

    pdf_content = export_pdf([], columns=_COLUMNS, title="Despesas")
    assert pdf_content[:5] == b"%PDF-"

    xlsx_content = export_xlsx([], columns=_COLUMNS, title="Despesas")
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_content))
    assert workbook.active is not None
